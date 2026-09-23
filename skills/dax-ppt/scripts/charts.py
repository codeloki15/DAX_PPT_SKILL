"""Structured chart specs and native PowerPoint chart insertion.

The agent never hand-codes chart geometry. It calls create_chart() with a
structured spec (type, categories, series); the spec is stored as JSON in the
workspace. The live preview renders it with Chart.js, and on export the same
spec becomes a NATIVE, editable PowerPoint chart (python-pptx add_chart) placed
exactly where the `.chart-embed` placeholder sat in the HTML slide.
"""

import json
import os
import re
from typing import Dict, Any, List, Optional

from paths import CHARTS_DIR, ensure_directories

try:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False


# Brand series ramp (Data Axle): brandBlue first, then supporting neutrals.
BRAND_SERIES_COLORS = ["00A0DC", "12263F", "6A7C90", "8FD3EE", "3C4456", "D4D9E0"]
WATERFALL_UP = "00A0DC"
WATERFALL_DOWN = "B3341F"
WATERFALL_TOTAL = "12263F"
AXIS_TEXT = "6A7C90"
GRIDLINE = "E4E8ED"

CHART_TYPES = {
    "column", "bar", "line", "area", "pie", "doughnut",
    "column_stacked", "bar_stacked", "waterfall",
}
SINGLE_SERIES_TYPES = {"pie", "doughnut", "waterfall"}

_CHART_ID_RE = re.compile(r"^[a-zA-Z0-9_\-]{1,40}$")


# =============================================================================
# Spec creation / storage
# =============================================================================

def _spec_path(chart_id: str) -> str:
    return os.path.join(CHARTS_DIR, f"{chart_id}.json")


def _normalize_color(value: str) -> Optional[str]:
    value = str(value).lstrip("#").upper()
    return value if re.fullmatch(r"[0-9A-F]{6}", value) else None


def create_chart(
    chart_id: str,
    chart_type: str,
    categories: List[str],
    series: List[Dict[str, Any]],
    title: Optional[str] = None,
    colors: Optional[List[str]] = None,
    number_format: Optional[str] = None,
    totals: Optional[List[int]] = None,
) -> Dict[str, Any]:
    """Validate and store a chart spec; return the HTML placeholder to embed."""
    ensure_directories()

    if not _CHART_ID_RE.fullmatch(str(chart_id)):
        return {"error": "chart_id must be 1-40 chars of letters, digits, '_' or '-'"}

    if chart_type not in CHART_TYPES:
        return {"error": f"Unknown chart_type '{chart_type}'. Valid: {sorted(CHART_TYPES)}"}

    if not categories or not isinstance(categories, list):
        return {"error": "categories must be a non-empty list of labels"}
    categories = [str(c) for c in categories]

    if not series or not isinstance(series, list):
        return {"error": "series must be a non-empty list of {name, values} objects"}

    if chart_type in SINGLE_SERIES_TYPES and len(series) != 1:
        return {"error": f"chart_type '{chart_type}' requires exactly one series"}

    clean_series = []
    for i, s in enumerate(series):
        if not isinstance(s, dict) or "values" not in s:
            return {"error": f"series[{i}] must be an object with 'name' and 'values'"}
        values = s["values"]
        if len(values) != len(categories):
            return {"error": f"series[{i}] has {len(values)} values but there are "
                             f"{len(categories)} categories - they must match"}
        try:
            values = [float(v) for v in values]
        except (TypeError, ValueError):
            return {"error": f"series[{i}] values must all be numbers"}
        clean_series.append({"name": str(s.get("name", f"Series {i + 1}")), "values": values})

    clean_colors = None
    if colors:
        clean_colors = [c for c in (_normalize_color(c) for c in colors) if c]
        if len(clean_colors) != len(colors):
            return {"error": "colors must be 6-digit hex values like '00A0DC'"}

    clean_totals = None
    if totals is not None:
        if chart_type != "waterfall":
            return {"error": "totals only applies to waterfall charts"}
        try:
            clean_totals = sorted({int(t) for t in totals})
        except (TypeError, ValueError):
            return {"error": "totals must be a list of category indices (integers)"}
        if clean_totals and (clean_totals[0] < 0 or clean_totals[-1] >= len(categories)):
            return {"error": f"totals indices must be within 0..{len(categories) - 1}"}

    spec = {
        "chart_id": chart_id,
        "type": chart_type,
        "title": str(title) if title else None,
        "categories": categories,
        "series": clean_series,
        "colors": clean_colors,
        "number_format": str(number_format) if number_format else None,
        "totals": clean_totals or [],
    }

    with open(_spec_path(chart_id), "w", encoding="utf-8") as f:
        json.dump(spec, f, indent=2)

    embed_html = f'<div class="chart-embed" data-chart-id="{chart_id}" style="flex:1;min-height:200px;"></div>'
    return {
        "status": "created",
        "chart_id": chart_id,
        "embed_html": embed_html,
        "usage": (
            "Paste embed_html into the slide inside a container that has real height, e.g. "
            '<div class="exh" style="flex:1;display:flex;flex-direction:column;">'
            '<div class="exh-t">EXHIBIT 1 | ...</div>' + embed_html + "</div>. "
            "The preview renders it live; the PPTX export replaces it with a native, "
            "editable PowerPoint chart at the same position."
        ),
    }


def load_chart_spec(chart_id: str) -> Optional[Dict[str, Any]]:
    path = _spec_path(chart_id)
    if not os.path.exists(path):
        return None
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def load_all_chart_specs() -> Dict[str, Dict[str, Any]]:
    """All stored specs keyed by chart_id (injected into the live preview)."""
    specs = {}
    if os.path.isdir(CHARTS_DIR):
        for name in sorted(os.listdir(CHARTS_DIR)):
            if name.endswith(".json"):
                try:
                    with open(os.path.join(CHARTS_DIR, name), "r", encoding="utf-8") as f:
                        spec = json.load(f)
                    specs[spec["chart_id"]] = spec
                except (json.JSONDecodeError, KeyError):
                    continue
    return specs


# =============================================================================
# Native chart insertion (post-processing the exported PPTX)
# =============================================================================

_XL_TYPE_MAP = {
    "column": "COLUMN_CLUSTERED",
    "bar": "BAR_CLUSTERED",
    "line": "LINE",
    "area": "AREA",
    "pie": "PIE",
    "doughnut": "DOUGHNUT",
    "column_stacked": "COLUMN_STACKED",
    "bar_stacked": "BAR_STACKED",
    "waterfall": "COLUMN_STACKED",  # classic consulting trick: hidden base series
}


def _waterfall_segments(values: List[float], totals: List[int]):
    """Return (base, height, color) per bar for the stacked-column waterfall.

    Bars at `totals` indices are full bars from zero: a non-zero value anchors
    the running total (e.g. the FY24 opening bar), a zero value renders the
    computed running total (e.g. the FY25 closing bar).
    """
    totals_set = set(totals)
    bases, heights, colors = [], [], []
    running = 0.0
    for i, v in enumerate(values):
        if i in totals_set:
            if v:
                running = v
            bases.append(0.0)
            heights.append(running)
            colors.append(WATERFALL_TOTAL)
        elif v >= 0:
            bases.append(running)
            heights.append(v)
            colors.append(WATERFALL_UP)
            running += v
        else:
            bases.append(running + v)
            heights.append(-v)
            colors.append(WATERFALL_DOWN)
            running += v
    return bases, heights, colors


def _series_color(spec: Dict[str, Any], index: int) -> str:
    colors = spec.get("colors") or []
    if index < len(colors):
        return colors[index]
    return BRAND_SERIES_COLORS[index % len(BRAND_SERIES_COLORS)]


def _style_chart(chart, spec: Dict[str, Any]) -> None:
    """Best-effort brand styling; never allowed to break the export."""
    try:
        chart.font.size = Pt(9)
        chart.font.name = "Poppins"
        chart.font.color.rgb = RGBColor.from_string("3C4456")
    except Exception:
        pass

    try:
        if spec.get("title"):
            chart.has_title = True
            chart.chart_title.text_frame.text = spec["title"]
            run = chart.chart_title.text_frame.paragraphs[0].runs[0]
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string("12263F")
        else:
            chart.has_title = False
    except Exception:
        pass

    multi_series = len(spec["series"]) > 1 and spec["type"] != "waterfall"
    try:
        chart.has_legend = multi_series or spec["type"] in ("pie", "doughnut")
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
            chart.legend.font.size = Pt(9)
    except Exception:
        pass

    if spec["type"] not in ("pie", "doughnut"):
        try:
            cat_axis = chart.category_axis
            cat_axis.tick_labels.font.size = Pt(9)
            cat_axis.tick_labels.font.color.rgb = RGBColor.from_string(AXIS_TEXT)
            cat_axis.has_major_gridlines = False
        except Exception:
            pass
        try:
            val_axis = chart.value_axis
            val_axis.tick_labels.font.size = Pt(9)
            val_axis.tick_labels.font.color.rgb = RGBColor.from_string(AXIS_TEXT)
            val_axis.has_major_gridlines = True
            val_axis.major_gridlines.format.line.color.rgb = RGBColor.from_string(GRIDLINE)
            if spec.get("number_format"):
                val_axis.tick_labels.number_format = spec["number_format"]
                val_axis.tick_labels.number_format_is_linked = False
        except Exception:
            pass


def _color_series(chart, spec: Dict[str, Any]) -> None:
    try:
        plot = chart.plots[0]
        if spec["type"] in ("pie", "doughnut"):
            series = plot.series[0]
            for i, point in enumerate(series.points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = RGBColor.from_string(_series_color(spec, i))
        else:
            for i, series in enumerate(plot.series):
                color = RGBColor.from_string(_series_color(spec, i))
                if spec["type"] in ("line", "area"):
                    series.format.line.color.rgb = color
                    series.format.line.width = Pt(2)
                if spec["type"] != "line":
                    series.format.fill.solid()
                    series.format.fill.fore_color.rgb = color
    except Exception:
        pass


def _add_waterfall(slide, spec: Dict[str, Any], x, y, cx, cy):
    values = spec["series"][0]["values"]
    bases, heights, colors = _waterfall_segments(values, spec.get("totals") or [])

    chart_data = CategoryChartData()
    chart_data.categories = spec["categories"]
    chart_data.add_series("_base", bases)
    chart_data.add_series(spec["series"][0]["name"], heights)

    frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, x, y, cx, cy, chart_data)
    chart = frame.chart
    _style_chart(chart, spec)

    try:
        plot = chart.plots[0]
        plot.gap_width = 60
        base_series, value_series = plot.series[0], plot.series[1]
        base_series.format.fill.background()
        base_series.format.line.fill.background()
        for i, point in enumerate(value_series.points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = RGBColor.from_string(colors[i])
    except Exception:
        pass
    return chart


def _add_standard_chart(slide, spec: Dict[str, Any], x, y, cx, cy):
    chart_data = CategoryChartData()
    chart_data.categories = spec["categories"]
    for s in spec["series"]:
        chart_data.add_series(s["name"], s["values"])

    xl_type = getattr(XL_CHART_TYPE, _XL_TYPE_MAP[spec["type"]])
    frame = slide.shapes.add_chart(xl_type, x, y, cx, cy, chart_data)
    chart = frame.chart
    _style_chart(chart, spec)
    _color_series(chart, spec)

    # Value labels on simple single-series bars keep the consulting look.
    if spec["type"] in ("column", "bar") and len(spec["series"]) == 1:
        try:
            plot = chart.plots[0]
            plot.has_data_labels = True
            plot.data_labels.font.size = Pt(9)
            plot.data_labels.font.color.rgb = RGBColor.from_string("221F20")
            if spec.get("number_format"):
                plot.data_labels.number_format = spec["number_format"]
                plot.data_labels.number_format_is_linked = False
        except Exception:
            pass
    return chart


def insert_native_tables(pptx_path: str, placements: List[Dict[str, Any]]) -> Dict[str, Any]:
    """Insert native PowerPoint tables where HTML <table> exhibits sat.

    dom-to-pptx rasterises an HTML table into a single flat picture, so the
    numbers stop being selectable or editable. Rebuilding them as real PPTX
    tables keeps every cell as text the recipient can edit.

    placements: [{dom_index, x, y, w, h, headers: [...], rows: [[...]],
                  aligns: ["l"|"r", ...]}] with x/y/w/h as slide fractions.
    """
    if not PPTX_AVAILABLE:
        return {"inserted": 0, "errors": ["python-pptx not installed"]}
    if not placements:
        return {"inserted": 0, "errors": []}

    prs = Presentation(pptx_path)
    slides = list(prs.slides)
    slide_w, slide_h = prs.slide_width, prs.slide_height

    inserted, errors = 0, []
    for p in placements:
        idx = p.get("dom_index", -1)
        headers = p.get("headers") or []
        rows = p.get("rows") or []
        if not 0 <= idx < len(slides) or not headers or not rows:
            continue
        try:
            n_rows, n_cols = len(rows) + 1, len(headers)
            x = Emu(int(p["x"] * slide_w))
            y = Emu(int(p["y"] * slide_h))
            cx = Emu(int(p["w"] * slide_w))
            cy = Emu(int(p["h"] * slide_h))

            shape = slides[idx].shapes.add_table(n_rows, n_cols, x, y, cx, cy)
            table = shape.table
            aligns = p.get("aligns") or []

            def _style(cell, text, *, header):
                cell.text = str(text)
                para = cell.text_frame.paragraphs[0]
                if aligns and len(aligns) == n_cols:
                    col = cell._tc.getparent().index(cell._tc)
                    if col < len(aligns) and aligns[col] == "r":
                        para.alignment = PP_ALIGN.RIGHT
                for run in para.runs:
                    run.font.size = Pt(9 if header else 10)
                    run.font.name = "Poppins"
                    run.font.bold = header
                    run.font.color.rgb = RGBColor.from_string(
                        "6A7C90" if header else "3C4456")
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string(
                    "FFFFFF" if header else "FFFFFF")

            for c, head in enumerate(headers):
                _style(table.cell(0, c), head, header=True)
            for r, row in enumerate(rows, start=1):
                for c in range(n_cols):
                    _style(table.cell(r, c), row[c] if c < len(row) else "", header=False)

            inserted += 1
        except Exception as e:
            errors.append(f"slide {idx + 1}: {e}")

    if inserted:
        prs.save(pptx_path)
    return {"inserted": inserted, "errors": errors}


def insert_speaker_notes(pptx_path: str, notes_by_index: Dict[int, str]) -> Dict[str, Any]:
    """Write speaker notes into the exported deck's native notes fields.

    notes_by_index maps 0-based slide order to note text.
    """
    if not PPTX_AVAILABLE:
        return {"written": 0, "errors": ["python-pptx not installed"]}
    if not notes_by_index:
        return {"written": 0, "errors": []}

    prs = Presentation(pptx_path)
    slides = list(prs.slides)
    written, errors = 0, []
    for idx, text in notes_by_index.items():
        if not text or not 0 <= idx < len(slides):
            continue
        try:
            slides[idx].notes_slide.notes_text_frame.text = text
            written += 1
        except Exception as e:
            errors.append(f"slide {idx + 1}: {e}")

    if written:
        prs.save(pptx_path)
    return {"written": written, "errors": errors}


def insert_native_charts(pptx_path: str, placements: List[Dict[str, Any]]) -> Dict[str, Any]:
    """Insert native charts into an exported PPTX.

    placements: [{dom_index, chart_id, x, y, w, h}] where x/y/w/h are fractions
    of the slide (0..1) measured from the HTML placeholder's bounding box, and
    dom_index is the 0-based slide order in the exported deck.
    """
    if not PPTX_AVAILABLE:
        return {"inserted": [], "errors": ["python-pptx not installed. Run: pip install python-pptx"]}

    prs = Presentation(pptx_path)
    slides = list(prs.slides)
    slide_w, slide_h = prs.slide_width, prs.slide_height

    inserted, errors = [], []
    for p in placements:
        chart_id = p.get("chart_id", "?")
        idx = p.get("dom_index", -1)
        if not 0 <= idx < len(slides):
            errors.append(f"{chart_id}: slide index {idx} out of range (deck has {len(slides)})")
            continue
        spec = load_chart_spec(chart_id)
        if spec is None:
            errors.append(f"{chart_id}: no stored spec found")
            continue
        try:
            x = Emu(int(p["x"] * slide_w))
            y = Emu(int(p["y"] * slide_h))
            cx = Emu(max(int(p["w"] * slide_w), Emu(914400)))   # at least 1 inch
            cy = Emu(max(int(p["h"] * slide_h), Emu(457200)))   # at least 0.5 inch
            if spec["type"] == "waterfall":
                _add_waterfall(slides[idx], spec, x, y, cx, cy)
            else:
                _add_standard_chart(slides[idx], spec, x, y, cx, cy)
            inserted.append(chart_id)
        except Exception as e:
            errors.append(f"{chart_id}: {e}")

    if inserted:
        prs.save(pptx_path)
    return {"inserted": inserted, "errors": errors}
